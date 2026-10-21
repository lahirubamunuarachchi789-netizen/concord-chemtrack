"""Reusable UI building blocks (HTML cards, headers, KPI tiles, tables)."""

from __future__ import annotations

import html
from pathlib import Path
from typing import Iterable, Optional

import pandas as pd
import streamlit as st

APP_NAME = "Concord ChemTrack"
COMPANY = "Concord Footwear (Pvt) Ltd"


def esc(value: object) -> str:
    """Escape user-provided strings before embedding them in HTML."""
    return html.escape("" if value is None else str(value))


# --------------------------------------------------------------------------
# Branding / headers
# --------------------------------------------------------------------------
def brand_block(compact: bool = False) -> None:
    """Logo + product name, used in the sidebar and on the auth screen."""
    initials = "CF"
    sub = "" if compact else '<div class="cc-brand-sub">Chemical Inventory</div>'
    st.markdown(
        f"""
        <div class="cc-brand">
          <div class="cc-brand-mark">{initials}</div>
          <div>
            <div class="cc-brand-name">Concord <span style="font-weight:500">ChemTrack</span></div>
            {sub}
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def page_header(title: str, subtitle: str = "", right_html: str = "") -> None:
    """Standard page header bar with an optional right-hand slot."""
    st.markdown(
        f"""
        <div class="cc-page-header">
          <div>
            <div class="cc-title">{esc(title)}</div>
            <div class="cc-sub">{esc(subtitle)}</div>
          </div>
          <div>{right_html}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def user_chip(username: str, role: str = "Inventory User") -> str:
    """Return HTML for the small user pill shown in the sidebar/header."""
    initial = esc(username[:1].upper()) if username else "?"
    return f"""
    <div class="cc-user">
      <div class="cc-avatar">{initial}</div>
      <div style="min-width:0">
        <div class="cc-user-name">{esc(username)}</div>
        <div class="cc-user-role">{esc(role)}</div>
      </div>
    </div>
    """


# --------------------------------------------------------------------------
# KPI tiles & cards
# --------------------------------------------------------------------------
def kpi(label: str, value: object, foot: str = "") -> None:
    """Render one KPI tile (works inside ``st.columns``)."""
    st.markdown(
        f"""
        <div class="cc-kpi">
          <div class="cc-kpi-label">{esc(label)}</div>
          <div class="cc-kpi-value">{esc(value)}</div>
          <div class="cc-kpi-foot">{esc(foot)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def section_title(title: str, hint: str = "") -> None:
    hint_html = f'<span class="cc-muted" style="font-size:.78rem">{esc(hint)}</span>' if hint else ""
    st.markdown(
        f"""
        <div style="display:flex;justify-content:space-between;align-items:baseline;
                    gap:.5rem;margin:.35rem 0 .5rem 0;">
          <h3 style="margin:0">{esc(title)}</h3>{hint_html}
        </div>
        """,
        unsafe_allow_html=True,
    )


def card(title: str, body_html: str) -> None:
    st.markdown(
        f'<div class="cc-card"><h4>{esc(title)}</h4>{body_html}</div>',
        unsafe_allow_html=True,
    )


def setup_card(error: Optional[str], secrets_path: Path) -> None:
    """Shown when Supabase credentials are missing / invalid."""
    error_html = (
        f'<div style="margin-top:.5rem;color:#A62B2B;font-size:.82rem">Last error: {esc(error)}</div>'
        if error else ""
    )
    st.markdown(
        f"""
        <div class="cc-auth-wrap">
          <div class="cc-auth-brand">
            <div class="cc-auth-logo">CF</div>
            <div class="cc-auth-title">Setup required</div>
            <div class="cc-auth-sub">Connect {esc(APP_NAME)} to your Supabase project</div>
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.warning(
        "Supabase credentials were not found or still contain placeholders.\n\n"
        f"1. Open `{secrets_path}`\n"
        "2. Set `SUPABASE_URL` and `SUPABASE_KEY` (Supabase → Project Settings → API)\n"
        "3. Save and refresh this page."
        + (f"\n\n**Details:** {error}" if error else ""),
        icon="🔐",
    )
    with st.expander("Show required database schema (SQL)"):
        from src.config import SCHEMA_SQL

        st.code(SCHEMA_SQL, language="sql")


# --------------------------------------------------------------------------
# Tables
# --------------------------------------------------------------------------
def dataframe_table(df: pd.DataFrame, height: Optional[int] = None, column_config=None) -> None:
    """Render a consistent, responsive table, or a friendly empty state."""
    if df is None or df.empty:
        st.info("No records to display for the current selection.", icon="📄")
        return
    st.dataframe(
        df,
        width="stretch",
        hide_index=True,
        height=height,
        column_config=column_config,
    )


def download_excel(df: pd.DataFrame, filename: str) -> bytes:
    """Build an in-memory .xlsx workbook from a dataframe."""
    import io

    from openpyxl.utils import get_column_letter

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Report")
        sheet = writer.sheets["Report"]
        for idx, column in enumerate(df.columns, start=1):
            values = df[column].astype(str).head(200)
            longest = max([len(str(column)), *(len(v) for v in values.tolist())] or [10])
            sheet.column_dimensions[get_column_letter(idx)].width = max(12, min(34, longest + 3))
    return buffer.getvalue()


def sortable_columns(columns: Iterable[str]) -> list[str]:
    return [c for c in columns if c]
